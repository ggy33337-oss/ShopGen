import base64
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path


SUPPORTED_FILE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".pdf", ".docx"}
MAX_EXTRACTED_TEXT_LENGTH = 12000
MAX_EXTRACTED_IMAGES = 6
MAX_IMAGE_BYTES = 4 * 1024 * 1024


@dataclass(frozen=True)
class UploadedContent:
    filename: str
    content_type: str
    extension: str
    text: str = ""
    data_urls: tuple[str, ...] = ()

    @property
    def is_image(self):
        return self.extension in {".png", ".jpg", ".jpeg"}

    @property
    def data_url(self):
        if not self.data_urls:
            return ""
        return self.data_urls[0]


def normalize_extension(filename):
    return Path(str(filename or "")).suffix.lower()


def build_data_url(extension, content):
    media_type = get_media_type(extension)
    encoded = base64.b64encode(content).decode("ascii")
    return f"data:{media_type};base64,{encoded}"


def get_media_type(extension):
    extension = str(extension or "").lower()
    if extension == ".png":
        return "image/png"
    if extension in {".jpg", ".jpeg"}:
        return "image/jpeg"
    if extension == ".gif":
        return "image/gif"
    if extension == ".webp":
        return "image/webp"
    return "image/png"


def append_image_data_url(data_urls, extension, content):
    if not content or len(content) > MAX_IMAGE_BYTES:
        return
    if len(data_urls) >= MAX_EXTRACTED_IMAGES:
        return
    data_urls.append(build_data_url(extension, content))


def extract_pdf_text(content):
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise RuntimeError("缺少 PDF 解析依赖 pypdf，请先安装 requirements.txt。") from e

    reader = PdfReader(BytesIO(content))
    page_texts = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        if page_text.strip():
            page_texts.append(page_text.strip())
        if len("\n".join(page_texts)) >= MAX_EXTRACTED_TEXT_LENGTH:
            break
    return "\n".join(page_texts)[:MAX_EXTRACTED_TEXT_LENGTH]


def extract_pdf_images(content):
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise RuntimeError("缺少 PDF 解析依赖 pypdf，请先安装 requirements.txt。") from e

    reader = PdfReader(BytesIO(content))
    data_urls = []
    for page in reader.pages:
        for image in getattr(page, "images", []):
            extension = Path(str(getattr(image, "name", ""))).suffix.lower() or ".png"
            append_image_data_url(data_urls, extension, image.data)
            if len(data_urls) >= MAX_EXTRACTED_IMAGES:
                return tuple(data_urls)
    return tuple(data_urls)


def extract_docx_text(content):
    try:
        from docx import Document
    except ImportError as e:
        raise RuntimeError("缺少 DOCX 解析依赖 python-docx，请先安装 requirements.txt。") from e

    document = Document(BytesIO(content))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    return "\n".join(paragraphs)[:MAX_EXTRACTED_TEXT_LENGTH]


def extract_docx_images(content):
    try:
        from docx import Document
    except ImportError as e:
        raise RuntimeError("缺少 DOCX 解析依赖 python-docx，请先安装 requirements.txt。") from e

    document = Document(BytesIO(content))
    data_urls = []
    seen_relation_ids = set()
    for shape in document.inline_shapes:
        relation_id = getattr(shape._inline.graphic.graphicData.pic.blipFill.blip, "embed", None)
        if not relation_id or relation_id in seen_relation_ids:
            continue
        seen_relation_ids.add(relation_id)
        image_part = document.part.related_parts.get(relation_id)
        if not image_part:
            continue
        extension = Path(str(getattr(image_part, "partname", ""))).suffix.lower() or ".png"
        append_image_data_url(data_urls, extension, image_part.blob)
        if len(data_urls) >= MAX_EXTRACTED_IMAGES:
            break
    return tuple(data_urls)


def parse_uploaded_file(filename, content_type, content):
    if not content:
        return None

    extension = normalize_extension(filename)
    if extension not in SUPPORTED_FILE_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_FILE_EXTENSIONS))
        raise ValueError(f"不支持的文件类型：{extension or '未知'}。当前支持：{supported}")

    if extension in {".png", ".jpg", ".jpeg"}:
        return UploadedContent(
            filename=filename,
            content_type=content_type or "",
            extension=extension,
            data_urls=(build_data_url(extension, content),),
        )

    if extension == ".pdf":
        extracted_text = extract_pdf_text(content)
        data_urls = extract_pdf_images(content)
    else:
        extracted_text = extract_docx_text(content)
        data_urls = extract_docx_images(content)

    return UploadedContent(
        filename=filename,
        content_type=content_type or "",
        extension=extension,
        text=extracted_text,
        data_urls=data_urls,
    )
