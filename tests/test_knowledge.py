# -*- coding: utf-8 -*-

import unittest
from io import BytesIO
from types import SimpleNamespace

from docx import Document
from qdrant_client import QdrantClient

from knowledge.chunker import build_knowledge_chunks, estimate_tokens, split_knowledge_text
from knowledge.document_parser import parse_knowledge_file
from knowledge.models import KnowledgeElement, KnowledgeImageRecord
from knowledge.qdrant_knowledge_base import QdrantKnowledgeBase
from knowledge.quality import is_meaningful_knowledge_text
from services.knowledge_service import KnowledgeService


class FakeRetrievalGateway:
    dimension = 4

    def embed_text(self, text):
        value = str(text or "")
        if "蓝" in value or "运动鞋" in value:
            return [1.0, 0.0, 0.0, 0.0]
        if "咖啡" in value:
            return [0.0, 1.0, 0.0, 0.0]
        return [0.0, 0.0, 1.0, 0.0]

    def embed_image(self, image_data_url, description=""):
        del image_data_url
        return self.embed_text(description)


class FailingImageGateway(FakeRetrievalGateway):
    def embed_image(self, image_data_url, description=""):
        del image_data_url, description
        raise RuntimeError("image embedding failed")


class FakeReranker:
    def rerank(self, query, candidates, top_n=3):
        del query
        priorities = {
            "paragraph": 0.95,
            "image": 0.94,
            "semantic": 0.90,
            "heading": 0.75,
            "title": 0.70,
        }
        ranked = []
        for candidate in candidates:
            item = dict(candidate)
            item["vector_score"] = float(item.get("score") or 0.0)
            item["rerank_score"] = priorities.get(item.get("chunk_type"), 0.69)
            item["score"] = item["rerank_score"]
            ranked.append(item)
        ranked.sort(key=lambda item: item["rerank_score"], reverse=True)
        return ranked[:top_n]


class FixedScoreReranker:
    def __init__(self, scores):
        self.scores = scores

    def rerank(self, query, candidates, top_n=3):
        del query
        ranked = []
        for candidate in candidates:
            item = dict(candidate)
            item["vector_score"] = float(item.get("score") or 0.0)
            item["rerank_score"] = float(self.scores[item["chunk_id"]])
            item["score"] = item["rerank_score"]
            ranked.append(item)
        ranked.sort(key=lambda item: item["rerank_score"], reverse=True)
        return ranked[:top_n]


class PaginatedScrollClient:
    def collection_exists(self, collection_name):
        del collection_name
        return True

    def scroll(self, collection_name, limit, offset, with_payload, with_vectors):
        del limit, with_payload, with_vectors
        total = 600 if collection_name == "text" else 6
        start = int(offset or 0)
        end = min(start + 256, total)
        points = [
            SimpleNamespace(
                payload={
                    "source_id": "large-source",
                    "source_title": "大型资料",
                    "source_filename": "large.pdf",
                    "category": "测试",
                }
            )
            for _ in range(start, end)
        ]
        return points, (end if end < total else None)


class KnowledgeChunkerTests(unittest.TestCase):
    def test_creates_title_heading_and_paragraph_chunks(self):
        text = "# 产品卖点\n海洋蓝鞋面，适合日常跑步。\n\n# 材质\n透气网布和耐磨橡胶底。"

        chunks = split_knowledge_text("source-1", "轻量运动鞋", text, category="鞋服")

        self.assertEqual("title", chunks[0].chunk_type)
        self.assertEqual("资料标题：轻量运动鞋", chunks[0].content)
        self.assertTrue(any(chunk.chunk_type == "heading" for chunk in chunks))
        self.assertTrue(any(chunk.chunk_type == "paragraph" for chunk in chunks))
        self.assertTrue(any("产品卖点" in chunk.content for chunk in chunks))

    def test_semantic_chunks_respect_size_limit(self):
        text = "。".join(["蓝色运动鞋适合日常训练"] * 160)

        chunks = split_knowledge_text(
            "source-2",
            "运动鞋说明",
            text,
            max_tokens=80,
            overlap_tokens=12,
        )

        semantic_chunks = [chunk for chunk in chunks if chunk.chunk_type == "semantic"]
        self.assertGreater(len(semantic_chunks), 1)
        self.assertTrue(all(estimate_tokens(chunk.content) <= 100 for chunk in semantic_chunks))

    def test_merges_only_short_paragraphs_in_the_same_section(self):
        elements = (
            KnowledgeElement("paragraph", "第一段完整说明。", 1, "第一节"),
            KnowledgeElement("paragraph", "第二段补充说明。", 1, "第一节"),
            KnowledgeElement("paragraph", "下一页独立说明。", 2, "第一节"),
        )

        chunks = build_knowledge_chunks("source-3", "语义切分", elements)
        semantic_chunks = [chunk for chunk in chunks if chunk.chunk_type == "semantic"]
        page_two_chunks = [chunk for chunk in chunks if chunk.page_number == 2]

        self.assertEqual(1, len(semantic_chunks))
        self.assertIn("第一段完整说明。", semantic_chunks[0].content)
        self.assertIn("第二段补充说明。", semantic_chunks[0].content)
        self.assertEqual("paragraph", page_two_chunks[0].chunk_type)

    def test_docx_parser_preserves_heading_paragraph_and_table(self):
        buffer = BytesIO()
        document = Document()
        document.add_heading("产品说明", level=1)
        document.add_paragraph("蓝色运动鞋适合日常跑步。")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "属性"
        table.cell(0, 1).text = "内容"
        table.cell(1, 0).text = "颜色"
        table.cell(1, 1).text = "蓝色"
        document.save(buffer)

        parsed = parse_knowledge_file("sample.docx", "", buffer.getvalue())

        self.assertEqual(
            ["heading", "paragraph", "table"],
            [element.element_type for element in parsed.elements],
        )
        self.assertIn("| 属性 | 内容 |", parsed.elements[-1].text)

    def test_drops_numeric_only_heading_during_chunking(self):
        elements = (
            KnowledgeElement("heading", "1.00", 37, "1.00"),
            KnowledgeElement("paragraph", "品牌标准蓝为 #0057B8。", 37, "1.00"),
        )

        chunks = build_knowledge_chunks("source-4", "品牌规范", elements)

        self.assertFalse(
            any(chunk.chunk_type == "heading" and chunk.title == "1.00" for chunk in chunks)
        )
        self.assertTrue(any("#0057B8" in chunk.content for chunk in chunks))

    def test_quality_filter_keeps_meaningful_numeric_table_and_rejects_garbage(self):
        table = "| 温度/℃ | 200 | 300 |\n| --- | --- | --- |\n| 氨含量/% | 89.9 | 71.0 |"

        self.assertTrue(is_meaningful_knowledge_text(table, "反应条件表"))
        self.assertFalse(is_meaningful_knowledge_text("1.00", "1.00"))
        self.assertFalse(is_meaningful_knowledge_text("锟斤拷锟斤拷 000", "乱码"))


class KnowledgeServiceTests(unittest.TestCase):
    def test_file_size_limit_comes_from_config(self):
        service = KnowledgeService(
            values={
                "KNOWLEDGE_MAX_FILE_MB": "1",
                "KNOWLEDGE_STORAGE_DIR": "data/test-knowledge",
            },
            knowledge_base=object(),
        )

        with self.assertRaisesRegex(ValueError, "不能超过 1 MB"):
            service.upload("large.pdf", "application/pdf", b"0" * (1024 * 1024 + 1))


class QdrantKnowledgeBaseTests(unittest.TestCase):
    def setUp(self):
        self.knowledge_base = QdrantKnowledgeBase(
            values={
                "QWEN_EMBEDDING_DIMENSION": "4",
                "QDRANT_TEXT_COLLECTION": "test_text",
                "QDRANT_IMAGE_COLLECTION": "test_image",
                "KNOWLEDGE_EMBEDDING_WORKERS": "1",
            },
            client=QdrantClient(":memory:"),
            retrieval_gateway=FakeRetrievalGateway(),
            reranker=FakeReranker(),
        )

    def test_stores_text_and_image_in_separate_collections_and_searches_both(self):
        chunks = split_knowledge_text(
            "shoe-source",
            "蓝色运动鞋",
            "# 产品说明\n海洋蓝鞋面，透气网布，适合日常跑步。",
            category="鞋服",
            source_filename="shoe.docx",
        )
        self.knowledge_base.upsert_text_chunks(chunks)
        self.knowledge_base.upsert_image(
            source_id="image-source",
            image_data_url="data:image/png;base64,aW1hZ2U=",
            title="蓝色鞋商品图",
            category="鞋服",
            description="蓝色运动鞋白底商品图",
            source_filename="shoe.png",
            tags=["蓝色", "运动鞋"],
        )

        matches = self.knowledge_base.search_matches("蓝色运动鞋", limit=5, category="鞋服")

        self.assertTrue(any(item["chunk_type"] == "paragraph" for item in matches))
        self.assertTrue(any(item["chunk_type"] == "image" for item in matches))
        self.assertTrue(all(item["category"] == "鞋服" for item in matches))
        self.assertTrue(all(item["rerank_score"] >= 0.68 for item in matches))
        self.assertTrue(all("vector_score" in item for item in matches))

    def test_delete_source_removes_its_vectors(self):
        chunks = split_knowledge_text("coffee-source", "咖啡豆", "深烘焙咖啡豆。", category="食品")
        self.knowledge_base.upsert_text_chunks(chunks)

        self.knowledge_base.delete_source("coffee-source")

        matches = self.knowledge_base.search_matches("咖啡", limit=5, category="食品")
        self.assertEqual([], matches)

    def test_replace_source_keeps_old_vectors_when_image_embedding_fails(self):
        old_chunks = split_knowledge_text(
            "stable-source",
            "旧资料",
            "旧资料内容。",
            category="测试",
        )
        self.knowledge_base.upsert_text_chunks(old_chunks)
        self.knowledge_base.retrieval_gateway = FailingImageGateway()
        new_chunks = split_knowledge_text(
            "stable-source",
            "新资料",
            "新资料内容。",
            category="测试",
        )
        image = KnowledgeImageRecord(
            source_id="stable-source",
            image_data_url="data:image/png;base64,aW1hZ2U=",
            title="测试图片",
            category="测试",
            description="测试图片",
            source_filename="test.png",
        )

        with self.assertRaisesRegex(RuntimeError, "image embedding failed"):
            self.knowledge_base.replace_source(new_chunks, [image])

        sources = self.knowledge_base.list_sources()
        self.assertEqual("旧资料", sources[0]["title"])

    def test_list_sources_counts_all_paginated_vectors(self):
        knowledge_base = QdrantKnowledgeBase(
            values={
                "QWEN_EMBEDDING_DIMENSION": "4",
                "QDRANT_TEXT_COLLECTION": "text",
                "QDRANT_IMAGE_COLLECTION": "image",
            },
            client=PaginatedScrollClient(),
            retrieval_gateway=FakeRetrievalGateway(),
        )

        sources = knowledge_base.list_sources()

        self.assertEqual(606, sources[0]["vector_count"])

    def test_filters_low_information_reranks_top_three_and_applies_threshold(self):
        reranker = FixedScoreReranker(
            {
                "direct": 0.91,
                "supporting": 0.72,
                "below-threshold": 0.54,
                "irrelevant": 0.30,
            }
        )
        knowledge_base = QdrantKnowledgeBase(
            values={
                "QWEN_EMBEDDING_DIMENSION": "4",
                "KNOWLEDGE_RECALL_LIMIT": "20",
                "KNOWLEDGE_RERANK_TOP_N": "3",
                "KNOWLEDGE_RERANK_MIN_SCORE": "0.55",
            },
            client=object(),
            retrieval_gateway=FakeRetrievalGateway(),
            reranker=reranker,
        )
        candidates = [
            make_candidate("direct", "品牌标准蓝为 #0057B8。", 0.80),
            make_candidate("supporting", "蓝色背景应保持纯净。", 0.79),
            make_candidate("below-threshold", "蓝色是一种冷色调。", 0.78),
            make_candidate("irrelevant", "商品支持顺丰发货。", 0.77),
            make_candidate("numeric", "1.00", 0.95, title="1.00"),
        ]
        knowledge_base.retrieve_candidates = lambda query, limit, category: candidates

        matches = knowledge_base.search_matches("背景改成品牌蓝色", limit=5)

        self.assertEqual(["direct", "supporting"], [item["chunk_id"] for item in matches])
        self.assertTrue(all(item["rerank_score"] >= 0.55 for item in matches))
        self.assertNotIn("numeric", [item["chunk_id"] for item in matches])

    def test_returns_no_match_when_every_rerank_score_is_below_threshold(self):
        knowledge_base = QdrantKnowledgeBase(
            values={
                "QWEN_EMBEDDING_DIMENSION": "4",
                "KNOWLEDGE_RERANK_MIN_SCORE": "0.55",
            },
            client=object(),
            retrieval_gateway=FakeRetrievalGateway(),
            reranker=FixedScoreReranker({"irrelevant": 0.42}),
        )
        knowledge_base.retrieve_candidates = lambda query, limit, category: [
            make_candidate("irrelevant", "商品支持顺丰发货。", 0.88)
        ]

        context = knowledge_base.search("背景改成品牌蓝色", limit=1)

        self.assertEqual("no_match", context.status)
        self.assertEqual((), context.matches)


def make_candidate(chunk_id, content, score, title="测试资料"):
    return {
        "source_id": "source",
        "chunk_id": chunk_id,
        "title": title,
        "content": content,
        "chunk_type": "paragraph",
        "category": "测试",
        "source_filename": "test.docx",
        "source_title": "测试资料",
        "page_number": 1,
        "section": "",
        "score": score,
    }


if __name__ == "__main__":
    unittest.main()
